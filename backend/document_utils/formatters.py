import io
import os
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from fpdf import FPDF


# =========================================================
# Logo
# =========================================================

PROJECT_ROOT = Path(
    __file__
).resolve().parents[2]


DEFAULT_LOGO = (
    PROJECT_ROOT
    / "assets"
    / "logo.png"
)


# =========================================================
# Text sanitization
# =========================================================

def sanitize_text(text: str) -> str:

    replacements = {

        "\u2018": "'",

        "\u2019": "'",

        "\u201c": '"',

        "\u201d": '"',

        "\u2013": "-",

        "\u2014": "-",

        "\u2026": "...",

        "\u00a0": " ",
    }


    for old, new in replacements.items():

        text = text.replace(
            old,
            new,
        )


    return text.strip()


# =========================================================
# TXT formatter
# =========================================================

def format_txt(
    text: str,
) -> bytes:

    clean_text = sanitize_text(
        text
    )

    return clean_text.encode(
        "utf-8"
    )


# =========================================================
# DOCX formatter
# =========================================================

def format_docx(
    text: str,
    document_type: str,
    terms: str = "",
    logo_bytes: bytes | None = None,
) -> bytes:

    document = Document()


    # -----------------------------------------------------
    # Page setup
    # -----------------------------------------------------

    section = document.sections[0]

    section.top_margin = Inches(
        0.8
    )

    section.bottom_margin = Inches(
        0.8
    )

    section.left_margin = Inches(
        0.9
    )

    section.right_margin = Inches(
        0.9
    )


    # -----------------------------------------------------
    # Header
    # -----------------------------------------------------

    header = section.header

    header_paragraph = (
        header.paragraphs[0]
    )

    header_paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )


    if logo_bytes:

        try:

            header_paragraph.add_run().add_picture(
                io.BytesIO(logo_bytes),
                width=Inches(1.2),
            )

        except Exception:

            pass


    header_run = (
        header_paragraph.add_run(
            f"\n{document_type.upper()}"
        )
    )

    header_run.bold = True

    header_run.font.name = (
        "Times New Roman"
    )

    header_run.font.size = Pt(11)


    # -----------------------------------------------------
    # Title
    # -----------------------------------------------------

    title = document.add_paragraph()

    title.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )


    title_run = title.add_run(
        document_type
    )

    title_run.bold = True

    title_run.font.name = (
        "Times New Roman"
    )

    title_run.font.size = Pt(18)


    # -----------------------------------------------------
    # Body
    # -----------------------------------------------------

    clean_text = sanitize_text(
        text
    )


    for line in clean_text.splitlines():

        line = line.strip()


        if not line:

            continue


        paragraph = (
            document.add_paragraph()
        )


        paragraph.paragraph_format.space_after = (
            Pt(6)
        )


        paragraph.paragraph_format.line_spacing = (
            1.15
        )


        run = paragraph.add_run(
            line
        )


        run.font.name = (
            "Times New Roman"
        )

        run.font.size = Pt(11)


        # Detect headings.
        is_heading = (
            line.isupper()
            or line.endswith(":")
            or (
                len(line) < 100
                and line[:2].isdigit()
                and "." in line[:5]
            )
        )


        if is_heading:

            run.bold = True


    # -----------------------------------------------------
    # Terms table
    # -----------------------------------------------------

    if terms.strip():

        heading = (
            document.add_paragraph()
        )

        heading_run = heading.add_run(
            "Key Terms"
        )

        heading_run.bold = True

        heading_run.font.name = (
            "Times New Roman"
        )

        heading_run.font.size = Pt(12)


        term_list = [

            term.strip()

            for term in terms.split(";")

            if term.strip()
        ]


        if term_list:

            table = document.add_table(
                rows=1,
                cols=2,
            )


            table.style = (
                "Table Grid"
            )


            table.rows[0].cells[0].text = (
                "No."
            )

            table.rows[0].cells[1].text = (
                "Term"
            )


            for index, term in enumerate(
                term_list,
                start=1,
            ):

                row = table.add_row()

                row.cells[0].text = (
                    str(index)
                )

                row.cells[1].text = (
                    term
                )


    # -----------------------------------------------------
    # Footer
    # -----------------------------------------------------

    footer = (
        section.footer.paragraphs[0]
    )

    footer.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )


    footer_run = footer.add_run(
        "LegalEase | AI-generated draft | "
        "Professional legal review recommended"
    )


    footer_run.font.name = (
        "Times New Roman"
    )

    footer_run.font.size = Pt(8)


    # -----------------------------------------------------
    # Save
    # -----------------------------------------------------

    output = io.BytesIO()

    document.save(
        output
    )


    return output.getvalue()


# =========================================================
# PDF class
# =========================================================

class LegalEasePDF(FPDF):

    def __init__(
        self,
        document_type: str,
        logo_path: str | None = None,
    ):

        super().__init__()

        self.document_type = (
            document_type
        )

        self.logo_path = logo_path


    def header(self):

        if (
            self.logo_path
            and os.path.exists(
                self.logo_path
            )
        ):

            try:

                self.image(
                    self.logo_path,
                    x=95,
                    y=8,
                    w=20,
                )

                self.ln(18)

            except Exception:

                pass


        self.set_font(
            "Times",
            "B",
            11,
        )


        self.cell(
            0,
            8,
            self.document_type.upper(),
            align="C",
        )


        self.ln(10)


    def footer(self):

        self.set_y(-15)


        self.set_font(
            "Times",
            "",
            8,
        )


        self.cell(
            0,
            10,
            "LegalEase | AI-generated draft | "
            "Review before use",
            align="C",
        )


# =========================================================
# PDF formatter
# =========================================================

def format_pdf(
    text: str,
    document_type: str,
    logo_path: str | None = None,
) -> bytes:

    if not logo_path:

        if DEFAULT_LOGO.exists():

            logo_path = str(
                DEFAULT_LOGO
            )


    pdf = LegalEasePDF(
        document_type=document_type,
        logo_path=logo_path,
    )


    pdf.set_auto_page_break(
        auto=True,
        margin=20,
    )


    pdf.add_page()


    pdf.set_font(
        "Times",
        "",
        11,
    )


    clean_text = sanitize_text(
        text
    )


    for line in clean_text.splitlines():

        line = line.strip()


        if not line:

            pdf.ln(3)

            continue


        is_heading = (
            line.isupper()
            or line.endswith(":")
            or (
                len(line) < 100
                and line[:2].isdigit()
                and "." in line[:5]
            )
        )


        if is_heading:

            pdf.set_font(
                "Times",
                "B",
                11,
            )

        else:

            pdf.set_font(
                "Times",
                "",
                11,
            )


        pdf.multi_cell(
            0,
            6,
            line,
        )


        pdf.ln(1)


    return bytes(
        pdf.output()
    )